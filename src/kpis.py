"""Executive KPI layer.

Computes the six leadership KPIs from the analytics mart and forecast results,
assembles them into an executive summary (value vs target vs status) plus
supporting by-category and monthly trend tables, and writes a formatted Excel
dashboard workbook.

    KPIs: Forecast Accuracy, Inventory Turnover, Fill Rate,
          Stockout %, Lead Time Variance, Supplier Performance

Outputs:
    data/processed/kpi_executive_summary.csv
    data/processed/kpi_by_category.csv
    data/processed/kpi_monthly.csv
    reports/executive_kpi_dashboard.xlsx

Run:  python src/kpis.py   (requires build_pipeline + forecasting to have run)
"""
from __future__ import annotations

import sqlite3

import numpy as np
import pandas as pd

from config import DB_PATH, PROCESSED_DIR, REPORTS_DIR

# KPI targets used to RAG-status the executive summary.
TARGETS = {
    "Forecast Accuracy (%)":        {"target": 85.0, "higher_better": True},
    "Inventory Turnover (x/yr)":    {"target": 8.0,  "higher_better": True},
    "Fill Rate (%)":                {"target": 95.0, "higher_better": True},
    "Stockout Rate (%)":            {"target": 2.0,  "higher_better": False},
    "Lead Time Variance (days)":    {"target": 2.0,  "higher_better": False},
    "Supplier Performance (score)": {"target": 80.0, "higher_better": True},
}


def _status(metric: str, value: float) -> str:
    t = TARGETS[metric]
    if t["higher_better"]:
        return "On Target" if value >= t["target"] else "Below Target"
    return "On Target" if value <= t["target"] else "Above Target"


def compute() -> dict:
    conn = sqlite3.connect(DB_PATH)
    sales = pd.read_sql(
        "SELECT date_key, category, quantity, revenue, cogs FROM fact_sales",
        conn, parse_dates=["date_key"])
    inv = pd.read_sql("SELECT * FROM fact_inventory", conn,
                      parse_dates=["snapshot_date"])
    pos = pd.read_sql("SELECT * FROM fact_purchase_orders", conn,
                      parse_dates=["received_date"])
    supplier_perf = pd.read_sql("SELECT * FROM v_supplier_performance", conn)
    demand_var = pd.read_sql("SELECT * FROM v_demand_variability", conn)
    conn.close()

    fc = pd.read_csv(REPORTS_DIR / "forecast_accuracy_summary.csv")
    fc_ml = fc[fc["model"].str.startswith("Gradient")]

    max_date = sales["date_key"].max()
    last_365 = max_date - pd.Timedelta(days=365)
    recent_sales = sales[sales["date_key"] > last_365]

    # --- KPI 1: Forecast Accuracy --------------------------------------
    forecast_acc = 100 - fc_ml.loc[fc_ml["category"] == "ALL", "MAPE"].iloc[0]

    # --- KPI 2: Inventory Turnover = annual COGS / avg inventory value --
    annual_cogs = recent_sales["cogs"].sum()
    avg_inv_value = (inv.groupby("snapshot_date")["inventory_value"].sum().mean())
    inv_turnover = annual_cogs / avg_inv_value

    # --- KPI 3: Fill Rate (units received / ordered) -------------------
    fill_rate = pos["fill_ratio"].mean() * 100

    # --- KPI 4: Stockout Rate ------------------------------------------
    stockout_rate = inv["is_stockout"].mean() * 100

    # --- KPI 5: Lead Time Variance -------------------------------------
    lead_time_var = pos["lead_time_variance_days"].mean()

    # --- KPI 6: Supplier Performance -----------------------------------
    supplier_score = supplier_perf["supplier_score"].mean()

    kpi_values = {
        "Forecast Accuracy (%)": round(forecast_acc, 1),
        "Inventory Turnover (x/yr)": round(inv_turnover, 1),
        "Fill Rate (%)": round(fill_rate, 1),
        "Stockout Rate (%)": round(stockout_rate, 2),
        "Lead Time Variance (days)": round(lead_time_var, 1),
        "Supplier Performance (score)": round(supplier_score, 1),
    }
    exec_summary = pd.DataFrame([
        {"KPI": k, "Value": v, "Target": TARGETS[k]["target"],
         "Status": _status(k, v)}
        for k, v in kpi_values.items()
    ])

    # --- By-category KPI table -----------------------------------------
    cat_cogs = recent_sales.groupby("category")["cogs"].sum()
    cat_inv = (inv.groupby(["snapshot_date", "category"])["inventory_value"]
                  .sum().groupby("category").mean())
    cat_turnover = (cat_cogs / cat_inv).round(1)
    cat_stockout = (inv.groupby("category")["is_stockout"].mean() * 100).round(2)
    cat_cv = demand_var.groupby("category")["cv_demand"].mean().round(3)
    cat_acc = (fc_ml[fc_ml["category"] != "ALL"]
               .assign(acc=lambda d: 100 - d["MAPE"])
               .set_index("category")["acc"].round(1))
    cat_rev = recent_sales.groupby("category")["revenue"].sum().round(0)

    by_category = pd.DataFrame({
        "Annual Revenue": cat_rev,
        "Forecast Accuracy (%)": cat_acc,
        "Inventory Turnover (x/yr)": cat_turnover,
        "Stockout Rate (%)": cat_stockout,
        "Demand Variability (CV)": cat_cv,
    }).reset_index().rename(columns={"index": "category"})

    # --- Monthly KPI trend ---------------------------------------------
    sales["ym"] = sales["date_key"].dt.to_period("M").astype(str)
    inv["ym"] = inv["snapshot_date"].dt.to_period("M").astype(str)
    pos["ym"] = pos["received_date"].dt.to_period("M").astype(str)

    monthly = pd.DataFrame({
        "Revenue": sales.groupby("ym")["revenue"].sum().round(0),
        "Units": sales.groupby("ym")["quantity"].sum(),
        "Stockout Rate (%)": (inv.groupby("ym")["is_stockout"].mean() * 100).round(2),
        "Fill Rate (%)": (pos.groupby("ym")["fill_ratio"].mean() * 100).round(1),
        "Lead Time Variance (days)": pos.groupby("ym")["lead_time_variance_days"].mean().round(1),
    }).reset_index().rename(columns={"index": "year_month", "ym": "year_month"})

    # --- Persist CSVs ---------------------------------------------------
    exec_summary.to_csv(PROCESSED_DIR / "kpi_executive_summary.csv", index=False)
    by_category.to_csv(PROCESSED_DIR / "kpi_by_category.csv", index=False)
    monthly.to_csv(PROCESSED_DIR / "kpi_monthly.csv", index=False)

    # --- Excel dashboard workbook --------------------------------------
    _write_excel(exec_summary, by_category, monthly, supplier_perf)

    # --- Console report -------------------------------------------------
    print("\n" + "=" * 56)
    print("EXECUTIVE KPI SCORECARD")
    print("=" * 56)
    for r in exec_summary.itertuples(index=False):
        flag = "OK " if r.Status == "On Target" else "!! "
        print(f"  {flag}{r.KPI:<30} {r.Value:>8}  (target {r.Target})")
    print("=" * 56)

    return kpi_values


def _write_excel(exec_summary, by_category, monthly, supplier_perf) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    path = REPORTS_DIR / "executive_kpi_dashboard.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        exec_summary.to_excel(xl, sheet_name="Executive Summary", index=False, startrow=1)
        by_category.to_excel(xl, sheet_name="KPIs by Category", index=False)
        monthly.to_excel(xl, sheet_name="Monthly Trend", index=False)
        supplier_perf.to_excel(xl, sheet_name="Supplier Scorecard", index=False)

        wb = xl.book
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(bold=True, color="FFFFFF")
        green = PatternFill("solid", fgColor="C6EFCE")
        red = PatternFill("solid", fgColor="FFC7CE")

        # Title + RAG colouring on the Executive Summary sheet.
        ws = wb["Executive Summary"]
        ws["A1"] = "Supply Chain Executive KPI Scorecard"
        ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
        for col in range(1, exec_summary.shape[1] + 1):
            c = ws.cell(row=2, column=col)
            c.fill, c.font = header_fill, header_font
        # Status column is the 4th; colour by value.
        for row in range(3, 3 + len(exec_summary)):
            cell = ws.cell(row=row, column=4)
            cell.fill = green if cell.value == "On Target" else red

        # Bold + colour headers and auto-size columns on every sheet.
        for name in wb.sheetnames:
            ws = wb[name]
            header_row = 2 if name == "Executive Summary" else 1
            for cell in ws[header_row]:
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(horizontal="center")
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(width + 3, 40)


if __name__ == "__main__":
    compute()
